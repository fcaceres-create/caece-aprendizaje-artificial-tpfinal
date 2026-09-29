"""Genera el informe final: plantilla Markdown + números reales → informe_final.md → DOCX (y PDF).

La plantilla ``reports/informe_template.md`` contiene marcadores:
  - ``{{clave}}``          → número formateado leído de las tablas generadas por el pipeline.
  - ``{{table:nombre}}``   → tabla Markdown construida a partir de un CSV.
  - ``{{include:ruta}}``   → contenido de otro archivo Markdown (por ejemplo el decision log).
Así ningún número del informe se escribe a mano: si se vuelve a correr el pipeline, el informe
se actualiza solo.
"""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import pandas as pd

from . import config as cfg

TEMPLATE = cfg.REPORTS_DIR / "informe_template.md"
REPORT_MD = cfg.REPORTS_DIR / "informe_final.md"
REPORT_DOCX = cfg.REPORTS_DIR / "informe_final.docx"
REPORT_PDF = cfg.REPORTS_DIR / "informe_final.pdf"


# --- Formato numérico en español -------------------------------------------------------------

def f(x, d: int = 3) -> str:
    """Decimal con coma y miles con punto: 12345.678 → '12.345,678'."""
    s = f"{float(x):,.{d}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def i(x) -> str:
    return f"{int(round(float(x))):,}".replace(",", ".")


def pct(x, d: int = 1) -> str:
    """x en proporción (0,123) → '12,3 %'."""
    return f"{f(100 * float(x), d)} %"


def _t(name: str) -> pd.DataFrame:
    return pd.read_csv(cfg.TABLES_DIR / name)


def _j(name: str) -> dict:
    return json.loads((cfg.TABLES_DIR / name).read_text(encoding="utf-8"))


# --- Números -------------------------------------------------------------------------------------

def collect_numbers() -> dict[str, str]:
    n: dict[str, str] = {"fecha": date.today().strftime("%d/%m/%Y")}
    info = _j("data_info.json")
    n.update(train_raw=i(info["train_raw_rows"]), train_clean=i(info["train_clean_rows"]),
             test_rows=i(info["test_rows"]), overlap=i(info["overlap_rows_removed"]))
    summ = _t("data_summary.csv").set_index("dataset")
    n.update(pos_train=i(summ.loc["train_clean", "positives"]),
             rate_train=f(summ.loc["train_clean", "positive_rate_pct"], 2) + " %",
             rate_train_raw=f(summ.loc["train_raw", "positive_rate_pct"], 2) + " %",
             rate_test=f(summ.loc["test", "positive_rate_pct"], 2) + " %",
             pos_test=i(summ.loc["test", "positives"]))
    sc = _t("split_checks.csv")
    n.update(ks_p=f(sc.iloc[0]["p_value"], 2), ks_d=f(sc.iloc[0]["statistic"], 3),
             chi_month_p=f(sc.iloc[1]["p_value"], 2), chi_contact_p=f(sc.iloc[2]["p_value"], 2),
             chi_pout_p=f(sc.iloc[3]["p_value"], 2), chi_y_p=f(sc.iloc[4]["p_value"], 2))

    eda = _j("eda_key_numbers.json")
    n.update(pct_pdays_neg=f(eda["pdays"]["pct_pdays_minus1"], 1) + " %",
             n_pdays_neg=i(eda["pdays"]["pdays_minus1_total"]),
             conv_never=pct(eda["pdays"]["conv_never_contacted"]),
             conv_before=pct(eda["pdays"]["conv_contacted_before"]),
             conv_success=pct(eda["pdays"]["conv_poutcome_success"]),
             spearman_month=f(eda["month"]["spearman_volume_vs_rate"][0], 2),
             may_share=f(eda["month"]["may_share_pct"], 1) + " %", may_rate=pct(eda["month"]["may_rate"]),
             bal_skew=f(eda["outliers"]["balance_skew"], 1), bal_min=i(eda["outliers"]["balance_min"]),
             bal_max=i(eda["outliers"]["balance_max"]), bal_neg=f(eda["outliers"]["balance_pct_negative"], 1) + " %",
             bal_zero=f(eda["outliers"]["balance_pct_zero"], 1) + " %",
             bal_iqr=f(eda["outliers"]["balance_pct_outliers_iqr"], 1) + " %",
             camp_max=i(eda["outliers"]["campaign_max"]), camp_p99=i(eda["outliers"]["campaign_p99"]),
             conv_camp1=pct(eda["outliers"]["conv_campaign_1"]), conv_camp5=pct(eda["outliers"]["conv_campaign_ge_5"]),
             dur_auc=f(eda["duration"]["duration_auc_alone"], 3),
             dur_med_no=i(eda["duration"]["duration_median_no"]), dur_med_yes=i(eda["duration"]["duration_median_yes"]),
             dur_short_conv=pct(eda["duration"]["conv_short_calls_lt_60s"]),
             dur_zero=i(eda["duration"]["duration_zero_n"]))
    unk = _t("eda_unknown_analysis.csv")
    for var in ["contact", "poutcome", "education"]:
        r = unk[(unk.variable == var)].set_index("is_unknown")["conversion_rate"]
        n[f"unk_{var}_conv"] = pct(r[True])
        n[f"unk_{var}_rest"] = pct(r[False])
    prof = _t("data_profile.csv")
    pu = prof[prof.dataset == "train_clean"].set_index("column")["pct_unknown"]
    for var in ["job", "education", "contact", "poutcome"]:
        n[f"pct_unk_{var}"] = f(pu[var], 1) + " %"

    cv = _t("cv_results.csv").set_index("model")
    for m in cv.index:
        n[f"cv_{m}_pr"] = f(cv.loc[m, "pr_auc_mean"], 3)
        n[f"cv_{m}_roc"] = f(cv.loc[m, "roc_auc_mean"], 3)
    imb = _t("imbalance_results.csv")
    for _, r in imb.iterrows():
        n[f"imb_{r['model']}_{r['imbalance']}_pr"] = f(r["pr_auc_mean"], 3)
        n[f"imb_{r['model']}_{r['imbalance']}_brier"] = f(r["brier_mean"], 3)
    tun = _t("tuning_summary.csv").set_index("model")
    for m in tun.index:
        n[f"tune_{m}_default"] = f(tun.loc[m, "default_pr_auc"], 4)
        n[f"tune_{m}_best"] = f(tun.loc[m, "best_pr_auc"], 4)
        n[f"tune_{m}_trials"] = i(tun.loc[m, "n_trials"])
        n[f"tune_{m}_secs"] = i(tun.loc[m, "elapsed_s"])
    dm = _t("decision_matrix.csv").set_index("candidate")
    n["dm_winner"] = dm.index[0]
    n["dm_winner_score"] = f(dm.iloc[0]["weighted_score"], 3)
    n["dm_logreg_score"] = f(dm.loc["logreg_default", "weighted_score"], 3)

    meta = json.loads(cfg.METADATA_FILE.read_text(encoding="utf-8"))
    hp = meta["hyperparameters"]
    n.update(model_label=meta["model_label"], hp_n_estimators=i(hp["n_estimators"]),
             hp_lr=f(hp["learning_rate"], 4), hp_leaves=i(hp["num_leaves"]),
             hp_min_child=i(hp["min_child_samples"]), hp_subsample=f(hp["subsample"], 2),
             hp_colsample=f(hp["colsample_bytree"], 2),
             cv_final_pr=f(meta["cv_metrics_mean"]["pr_auc"], 3), cv_final_pr_sd=f(meta["cv_metrics_std"]["pr_auc"], 3),
             cv_final_roc=f(meta["cv_metrics_mean"]["roc_auc"], 3), cv_final_lift20=f(meta["cv_metrics_mean"]["lift@20"], 2))
    thr = _j("threshold_choice.json")
    n.update(thr=f(thr["profit_threshold"], 3), thr_f1=f(thr["f1_threshold"], 3),
             thr_contacted=f(thr["contacted_pct_at_threshold"], 1) + " %", thr_recall=f(thr["recall_at_threshold"], 3),
             thr_profit=i(thr["profit_at_threshold"]), thr_profit_all=i(thr["profit_contact_all"]),
             thr_f1_profit=i(thr["profit_at_f1_threshold"]),
             thr_f1_contacted=f(thr["contacted_pct_at_f1_threshold"], 1) + " %")
    leak = _t("leakage_cv.csv").set_index("feature_set")
    n.update(leak_cv_pre_roc=f(leak.loc["pre_contact", "roc_auc_mean"], 3),
             leak_cv_dur_roc=f(leak.loc["with_duration", "roc_auc_mean"], 3),
             leak_cv_pre_pr=f(leak.loc["pre_contact", "pr_auc_mean"], 3),
             leak_cv_dur_pr=f(leak.loc["with_duration", "pr_auc_mean"], 3))

    tm = _t("test_metrics.csv").set_index("model")
    for m in tm.index:
        for k in ["roc_auc", "pr_auc", "lift@10", "lift@20", "gain@20", "brier"]:
            n[f"test_{m}_{k.replace('@', '')}"] = f(tm.loc[m, k], 3)
        n[f"test_{m}_thr"] = f(tm.loc[m, "threshold_at_thr"], 3)
    fin = tm.loc["final"]
    n.update(test_prec=pct(fin["precision_at_thr"]), test_rec=pct(fin["recall_at_thr"]),
             test_f1=f(fin["f1_at_thr"], 3), test_contacted=f(fin["contacted_pct_at_thr"], 1) + " %")
    cm = pd.read_csv(cfg.TABLES_DIR / "test_confusion_matrix.csv", index_col=0).to_numpy()
    n.update(cm_tn=i(cm[0, 0]), cm_fp=i(cm[0, 1]), cm_fn=i(cm[1, 0]), cm_tp=i(cm[1, 1]))
    dec = _t("test_deciles.csv")
    n.update(dec1_conv=i(dec.loc[0, "conversions"]), dec1_rate=pct(dec.loc[0, "conversion_rate"]),
             dec1_lift=f(dec.loc[0, "lift"], 2), gain10=f(dec.loc[0, "cum_gain_pct"], 1) + " %",
             gain20=f(dec.loc[1, "cum_gain_pct"], 1) + " %", gain30=f(dec.loc[2, "cum_gain_pct"], 1) + " %")
    biz = _t("test_business_summary.csv")
    op = biz[biz.strategy.str.startswith("Umbral")].iloc[0]
    allc = biz[biz.strategy == "Contactar a todos"].iloc[0]
    n.update(biz_profit=i(op["profit"]), biz_all=i(allc["profit"]), biz_uplift=i(op["profit_vs_all"]),
             biz_uplift_pct=f(100 * op["profit_vs_all"] / allc["profit"], 1) + " %",
             biz_conv=i(op["conversions"]))
    sens = _t("test_sensitivity.csv").set_index("vc_ratio")
    for r in sens.index:
        k = int(r)
        n[f"sens{k}_model"] = i(sens.loc[r, "profit_model"])
        n[f"sens{k}_all"] = i(sens.loc[r, "profit_contact_all"])
        n[f"sens{k}_contacted"] = f(sens.loc[r, "contacted_pct"], 1) + " %"
    tvc = _t("test_vs_cv.csv").set_index("metric")
    n["gap_pr"] = f(tvc.loc["pr_auc", "test"] - tvc.loc["pr_auc", "cv_mean"], 3)

    shap_imp = _t("shap_importance.csv")
    for k in range(5):
        n[f"shap{k + 1}"] = shap_imp.iloc[k]["label"]
        n[f"shap{k + 1}_pct"] = f(shap_imp.iloc[k]["share_pct"], 1) + " %"
    contact_share = shap_imp[shap_imp.feature.isin(["contact", "contact_known"])]["share_pct"].sum()
    n["shap_contact_share"] = f(contact_share, 1) + " %"
    ex = _j("explain_summary.json")
    n["rules_roc"] = f(ex["rules_tree"]["test_roc_auc"], 3)
    seg = _t("segment_metrics.csv")
    cu = seg[(seg.segment_var == "contact") & (seg.segment == "unknown")].iloc[0]
    n.update(seg_unknown_roc=f(cu["roc_auc"], 2), seg_unknown_n=i(cu["n"]), seg_unknown_rate=pct(cu["conversion_rate"]))
    old = seg[(seg.segment_var == "age_group") & (seg.segment == "65+")].iloc[0]
    n.update(seg_65_roc=f(old["roc_auc"], 2), seg_65_rate=pct(old["conversion_rate"]))
    ep = pd.read_csv(cfg.TABLES_DIR / "error_profiles.csv", index_col=0)
    n.update(fn_balance=i(ep.loc["FN", "balance"]), tp_balance=i(ep.loc["TP", "balance"]),
             fn_camp=f(ep.loc["FN", "campaign"], 1), tp_camp=f(ep.loc["TP", "campaign"], 1),
             fn_prev=f(ep.loc["FN", "pct_previously_contacted"], 0) + " %",
             tp_prev=f(ep.loc["TP", "pct_previously_contacted"], 0) + " %")
    ecp = _t("error_categorical_profiles.csv")
    cu2 = ecp[(ecp.variable == "contact") & (ecp.category == "unknown")].iloc[0]
    my = ecp[(ecp.variable == "month") & (ecp.category == "may")].iloc[0]
    n.update(fn_unknown=f(cu2["pct_FN"], 0) + " %", tp_unknown=f(cu2["pct_TP"], 0) + " %",
             fn_may=f(my["pct_FN"], 0) + " %", tp_may=f(my["pct_TP"], 0) + " %")
    rs = _t("reasonableness.csv")
    n["reason_ok"] = f"{int(rs['ok'].sum())} de {len(rs)}"
    return n


# --- Tablas --------------------------------------------------------------------------------------

def _md_table(df: pd.DataFrame) -> str:
    cols = list(df.columns)
    lines = ["| " + " | ".join(map(str, cols)) + " |", "|" + "|".join(["---"] * len(cols)) + "|"]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(str(v) for v in r.values) + " |")
    return "\n".join(lines)


def build_tables() -> dict[str, str]:
    t: dict[str, str] = {}
    steps = _t("data_preparation_steps.csv")
    t["prep_steps"] = _md_table(pd.DataFrame({
        "#": steps["order"], "Paso": steps["step"], "Conjunto": steps["dataset"],
        "Filas antes": steps["rows_before"].map(i), "Filas después": steps["rows_after"].map(i),
        "Detalle": steps["reason"]}))
    cat = _t("feature_catalog.csv")
    t["features"] = _md_table(pd.DataFrame({"Variable": cat["feature"], "Origen": cat["source"],
                                            "Transformación": cat["transformation"],
                                            "Justificación": cat["justification"]}))
    cv = _t("cv_results.csv")
    t["cv"] = _md_table(pd.DataFrame({
        "Modelo": cv["label"], "Familia": cv["family"],
        "PR-AUC": [f"{f(a)} ± {f(b)}" for a, b in zip(cv["pr_auc_mean"], cv["pr_auc_std"])],
        "ROC-AUC": [f"{f(a)} ± {f(b)}" for a, b in zip(cv["roc_auc_mean"], cv["roc_auc_std"])],
        "Lift@20 %": [f"{f(a, 2)} ± {f(b, 2)}" for a, b in zip(cv["lift@20_mean"], cv["lift@20_std"])],
        "Brier": ["n/a" if pd.isna(a) else f(a) for a in cv["brier_mean"]],
        "Ajuste (s)": [f(a, 1) for a in cv["fit_time_s_mean"]]}))
    enc = _t("encoding_comparison.csv")
    t["encoding"] = _md_table(pd.DataFrame({
        "Modelo": enc["model"], "Codificación": enc["kind"].map({"tree_onehot": "One-hot", "tree_native": "Nativa"}),
        "PR-AUC": [f"{f(a)} ± {f(b)}" for a, b in zip(enc["pr_auc_mean"], enc["pr_auc_std"])],
        "ROC-AUC": [f(a) for a in enc["roc_auc_mean"]]}))
    imb = _t("imbalance_results.csv")
    names = {"none": "Sin tratamiento", "weights": "Pesos de clase", "smote": "SMOTE"}
    t["imbalance"] = _md_table(pd.DataFrame({
        "Modelo": imb["model"], "Estrategia": imb["imbalance"].map(names),
        "PR-AUC": [f"{f(a)} ± {f(b)}" for a, b in zip(imb["pr_auc_mean"], imb["pr_auc_std"])],
        "ROC-AUC": [f(a) for a in imb["roc_auc_mean"]], "Lift@20 %": [f(a, 2) for a in imb["lift@20_mean"]],
        "Brier": [f(a) for a in imb["brier_mean"]]}))
    tun = _t("tuning_summary.csv")
    t["tuning"] = _md_table(pd.DataFrame({
        "Modelo": tun["model"], "Trials": tun["n_trials"], "Tiempo (s)": tun["elapsed_s"].map(lambda x: f(x, 0)),
        "PR-AUC por defecto": tun["default_pr_auc"].map(lambda x: f(x, 4)),
        "PR-AUC tuneado": tun["best_pr_auc"].map(lambda x: f(x, 4)),
        "Mejores hiperparámetros": tun["best_params"].map(
            lambda s: ", ".join(f"{k}={f(v, 4) if isinstance(v, float) else v}" for k, v in json.loads(s).items()))}))
    dm = _t("decision_matrix.csv")
    t["decision"] = _md_table(pd.DataFrame({
        "Candidato": dm["candidate"], "PR-AUC": dm["pr_auc"].map(f), "Desvío PR-AUC": dm["pr_auc_std"].map(f),
        "Lift@20 %": dm["lift@20"].map(lambda x: f(x, 2)), "Interpretab. (1-5)": dm["interpretability"],
        "ms / 1.000 filas": dm["predict_ms_per_1k"].map(lambda x: f(x, 1)),
        "Puntaje ponderado": dm["weighted_score"].map(f)}))
    sens = _t("sensitivity_oof.csv")
    t["sens_oof"] = _md_table(pd.DataFrame({
        "V/C": sens["vc_ratio"], "Umbral teórico C/V": sens["theoretical_threshold"].map(f),
        "Umbral óptimo OOF": sens["best_threshold"].map(f), "% contactado": sens["contacted_pct"].map(lambda x: f(x, 1)),
        "Recall": sens["recall"].map(f), "Beneficio modelo": sens["profit_model"].map(i),
        "Beneficio contactar a todos": sens["profit_contact_all"].map(i)}))
    tm = _t("test_metrics.csv")
    t["test"] = _md_table(pd.DataFrame({
        "Modelo": tm["label"], "ROC-AUC": tm["roc_auc"].map(f), "PR-AUC": tm["pr_auc"].map(f),
        "Lift@10 %": tm["lift@10"].map(lambda x: f(x, 2)), "Lift@20 %": tm["lift@20"].map(lambda x: f(x, 2)),
        "Gain@20 %": tm["gain@20"].map(pct), "Brier": tm["brier"].map(f),
        "Umbral": tm["threshold_at_thr"].map(f), "Precision": tm["precision_at_thr"].map(f),
        "Recall": tm["recall_at_thr"].map(f), "F1": tm["f1_at_thr"].map(f)}))
    tvc = _t("test_vs_cv.csv")
    t["test_vs_cv"] = _md_table(pd.DataFrame({"Métrica": tvc["metric"],
                                              "CV (media ± desvío)": [f"{f(a)} ± {f(b)}" for a, b in zip(tvc["cv_mean"], tvc["cv_std"])],
                                              "Test": tvc["test"].map(f)}))
    dec = _t("test_deciles.csv")
    t["deciles"] = _md_table(pd.DataFrame({
        "Decil": dec["decile"], "Clientes": dec["n"], "Conversiones": dec["conversions"],
        "Tasa": dec["conversion_rate"].map(pct), "Lift": dec["lift"].map(lambda x: f(x, 2)),
        "% contactado acum.": dec["cum_pct_contacted"].map(lambda x: f(x, 0) + " %"),
        "Gain acumulado": dec["cum_gain_pct"].map(lambda x: f(x, 1) + " %"),
        "Lift acumulado": dec["cum_lift"].map(lambda x: f(x, 2))}))
    biz = _t("test_business_summary.csv")
    t["business"] = _md_table(pd.DataFrame({
        "Estrategia": biz["strategy"], "Contactos": biz["contacted"].map(i),
        "% contactado": biz["contacted_pct"].map(lambda x: f(x, 1) + " %"),
        "Conversiones": biz["conversions"].map(i), "% de conversiones capturadas": biz["gain_pct"].map(lambda x: f(x, 1) + " %"),
        "Beneficio (V=20, C=1)": biz["profit"].map(i), "vs contactar a todos": biz["profit_vs_all"].map(i)}))
    sens = _t("test_sensitivity.csv")
    t["sens_test"] = _md_table(pd.DataFrame({
        "V/C": sens["vc_ratio"].map(lambda x: str(int(x))), "Umbral (elegido en OOF)": sens["threshold_from_oof"].map(f),
        "% contactado": sens["contacted_pct"].map(lambda x: f(x, 1) + " %"),
        "% conversiones": sens["gain_pct"].map(lambda x: f(x, 1) + " %"),
        "Beneficio modelo": sens["profit_model"].map(i), "Beneficio contactar a todos": sens["profit_contact_all"].map(i),
        "Diferencia": sens["profit_vs_all"].map(i)}))
    lk = _t("test_leakage.csv")
    t["leakage"] = _md_table(pd.DataFrame({"Modelo": lk["label"], "ROC-AUC": lk["roc_auc"].map(f),
                                           "PR-AUC": lk["pr_auc"].map(f), "Lift@10 %": lk["lift@10"].map(lambda x: f(x, 2)),
                                           "Gain@20 %": lk["gain@20"].map(pct)}))
    si = _t("shap_importance.csv").head(10)
    t["shap"] = _md_table(pd.DataFrame({"Variable": si["label"], "Media |SHAP|": si["mean_abs_shap"].map(lambda x: f(x, 4)),
                                        "% de la importancia": si["share_pct"].map(lambda x: f(x, 1) + " %")}))
    lr = _t("logreg_coefficients.csv").head(12)
    t["logreg"] = _md_table(pd.DataFrame({"Variable": lr["feature"], "Coeficiente": lr["coef"].map(f),
                                          "Odds ratio": lr["odds_ratio"].map(lambda x: f(x, 2))}))
    rl = _t("rules_tree_leaves.csv")
    t["rules"] = _md_table(pd.DataFrame({"Regla": rl["rule"], "Clientes (train)": rl["n_train"].map(i),
                                         "% del train": rl["share_pct"].map(lambda x: f(x, 1) + " %"),
                                         "Tasa de conversión": rl["conversion_rate"].map(pct),
                                         "Lift": rl["lift"].map(lambda x: f(x, 2))}))
    seg = _t("segment_metrics.csv")
    seg = seg[seg["reliable"]]
    t["segments"] = _md_table(pd.DataFrame({
        "Variable": seg["segment_var"], "Segmento": seg["segment"], "Clientes": seg["n"].map(i),
        "Conversiones": seg["conversions"].map(i), "Tasa": seg["conversion_rate"].map(pct),
        "ROC-AUC": seg["roc_auc"].map(lambda x: f(x, 3)), "Recall al umbral": seg["recall_at_thr"].map(lambda x: f(x, 3))}))
    ep = pd.read_csv(cfg.TABLES_DIR / "error_profiles.csv")
    t["errors"] = _md_table(pd.DataFrame({
        "Grupo": ep["grupo"], "Clientes": ep["n"].map(i), "Edad media": ep["age"].map(lambda x: f(x, 1)),
        "Saldo medio": ep["balance"].map(i), "Contactos campaña": ep["campaign"].map(lambda x: f(x, 2)),
        "% contactados antes": ep["pct_previously_contacted"].map(lambda x: f(x, 1) + " %"),
        "Prob. media": ep["score"].map(f)}))
    rs = _t("reasonableness.csv")
    t["reasonableness"] = _md_table(pd.DataFrame({"Chequeo": rs["check"], "Valor": rs["value"].map(f),
                                                  "Referencia": rs["reference"],
                                                  "Cumple": rs["ok"].map({True: "Sí", False: "No"})}))
    return t


def render_markdown() -> str:
    text = TEMPLATE.read_text(encoding="utf-8")
    nums = collect_numbers()
    tabs = build_tables()

    def repl(m):
        key = m.group(1).strip()
        if key.startswith("table:"):
            return tabs[key[6:]]
        if key.startswith("include:"):
            inc = (cfg.ROOT / key[8:]).read_text(encoding="utf-8")
            # Baja un nivel los títulos del archivo incluido para que queden dentro del anexo.
            return re.sub(r"^(#+)", lambda h: "#" + h.group(1), inc, flags=re.M)
        if key not in nums:
            raise KeyError(f"Marcador sin valor en la plantilla: {key}")
        return nums[key]

    out = re.sub(r"\{\{(.+?)\}\}", repl, text)
    REPORT_MD.write_text(out, encoding="utf-8")
    return out


# --- DOCX ------------------------------------------------------------------------------------------

_STRUCTURAL = re.compile(r"^\s*(#|\||!\[|```|<!--|>|\$\$|[-*] |\d+\. )")


def _merge_soft_wraps(lines: list[str]) -> list[str]:
    """Une las líneas cortadas de un mismo párrafo o ítem de lista (como hace Markdown)."""
    out: list[str] = []
    in_code = False
    for line in lines:
        s = line.strip()
        if s.startswith("```"):
            in_code = not in_code
            out.append(line)
            continue
        if in_code or not s or _STRUCTURAL.match(line):
            out.append(line)
            continue
        prev = out[-1] if out else ""
        mergeable = prev.strip() and (not _STRUCTURAL.match(prev) or re.match(r"^\s*([-*] |\d+\. |>)", prev))
        if mergeable:
            out[-1] = prev.rstrip() + " " + s
        else:
            out.append(line)
    return out


def _add_runs(par, text: str):
    """Negritas (**x**), cursivas (*x*) y código (`x`) en línea."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2].replace("`", "")).bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def markdown_to_docx(md: str, out: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    lines = _merge_soft_wraps(md.splitlines())
    k = 0
    in_code = False
    code_buf: list[str] = []
    while k < len(lines):
        line = lines[k]
        s = line.strip()
        if s.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                code_buf = []
            in_code = not in_code
            k += 1
            continue
        if in_code:
            code_buf.append(line)
            k += 1
            continue
        if s == "<!-- pagebreak -->":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif s.startswith("<!-- center -->"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(p, s.replace("<!-- center -->", "").strip())
        elif s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            doc.add_heading(s[level:].strip(), level=min(level, 4) if level > 1 else 0 if k < 3 else 1)
        elif s.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
            path = (REPORT_MD.parent / m.group(2)).resolve()
            if path.exists():
                doc.add_picture(str(path), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(m.group(1))
                r.italic = True
                r.font.size = Pt(9)
        elif s.startswith("|"):
            rows = []
            while k < len(lines) and lines[k].strip().startswith("|"):
                cells = [c.strip() for c in lines[k].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                k += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row[:len(rows[0])]):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    _add_runs(par, val)
                    for run in par.runs:
                        run.font.size = Pt(8)
                        if ri == 0:
                            run.bold = True
            doc.add_paragraph()
            continue
        elif re.match(r"^(\s*)[-*] ", line):
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            _add_runs(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\s*\d+\. ", "", line))
        elif s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs(p, s.lstrip("> ").strip())
        elif s.startswith("$$"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(s.strip("$").replace("\\cdot", "·").replace("\\text{", "").replace("}", ""))
            r.italic = True
        elif s:
            p = doc.add_paragraph()
            _add_runs(p, s)
        k += 1
    doc.save(out)


def build_report() -> dict:
    md = render_markdown()
    markdown_to_docx(md, REPORT_DOCX)
    result = {"markdown": str(REPORT_MD.relative_to(cfg.ROOT)), "docx": str(REPORT_DOCX.relative_to(cfg.ROOT)),
              "pdf": None}
    try:
        from docx2pdf import convert  # requiere Microsoft Word (Windows/macOS)
        convert(str(REPORT_DOCX), str(REPORT_PDF))
        result["pdf"] = str(REPORT_PDF.relative_to(cfg.ROOT))
    except Exception as exc:  # noqa: BLE001
        result["pdf_error"] = f"No se generó el PDF (se necesita Word + docx2pdf): {exc}"
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return result


if __name__ == "__main__":
    build_report()
